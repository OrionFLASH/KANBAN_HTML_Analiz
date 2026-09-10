"""Тесты текстовых идентификаторов для Excel."""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.excel_format import format_sheet, prepare_excel_frame
from src.text_ids import format_id_as_text, format_text_id_columns, text_id_column_labels


def test_format_id_as_text_preserves_long_digits() -> None:
    assert format_id_as_text("12345678901234567890") == "12345678901234567890"
    assert format_id_as_text(123456789012345) == "123456789012345"
    assert format_id_as_text(1.23e18) == "1230000000000000000"
    assert format_id_as_text("1.234567890123456e+18") is not None
    assert format_id_as_text(None) is None
    assert format_id_as_text(float("nan")) is None
    assert format_id_as_text("12345.0") == "12345"


def test_text_id_column_labels() -> None:
    config = {
        "columns": {"client_id": "Идентификатор клиента", "lead_id": "ID ПрПр"},
        "output": {
            "snapshot_columns": {"client_id": "Идентификатор клиента"},
        },
    }
    labels = text_id_column_labels(config)
    assert "Идентификатор клиента" in labels
    assert "ID ПрПр" in labels


def test_prepare_excel_frame_client_id_as_text(tmp_path: Path) -> None:
    config = {
        "columns": {
            "client_id": "Идентификатор клиента",
            "work_start_date": "Дата начала работы",
        },
        "dates": {"dayfirst": True, "formats": ["%Y-%m-%d", "%d.%m.%Y"], "empty_values": [""]},
        "output": {
            "snapshot_columns": {
                "client_id": "Идентификатор клиента",
                "work_start_date": "Дата начала работы",
            },
            "column_labels": {"min_header_marker": "Мин", "max_header_marker": "Макс"},
            "excel_format": {
                "date_format": "YYYY-MM-DD",
                "float_format": "0.00",
                "int_format": "0",
                "min_column_width": 12,
                "max_column_width": 45,
                "sample_rows_for_width": 50,
                "colors": {"min": "C6EFCE", "max": "FFC7CE"},
            },
        },
        "team_files": {"output_columns": {}},
        "excel_theme": "green_red",
    }
    frame = pd.DataFrame(
        {
            "Идентификатор клиента": [1234567890123456789, "9876543210987654321"],
            "Дата начала работы": ["2026-09-01", "15.08.2026"],
        }
    )
    prepared = prepare_excel_frame(frame, config)
    assert isinstance(prepared.loc[0, "Идентификатор клиента"], str)
    assert prepared.loc[1, "Идентификатор клиента"] == "9876543210987654321"

    path = tmp_path / "ids.xlsx"
    prepared.to_excel(path, index=False)
    wb = load_workbook(path)
    ws = wb.active
    format_sheet(ws, config, sheet_key="leads")
    assert ws.cell(2, 1).number_format == "@"
    assert isinstance(ws.cell(2, 1).value, str)
    assert ws.cell(2, 2).number_format == "YYYY-MM-DD"


def test_format_text_id_columns_frame() -> None:
    config = {
        "columns": {"client_id": "Идентификатор клиента"},
        "output": {"snapshot_columns": {"client_id": "Идентификатор клиента"}},
    }
    frame = pd.DataFrame({"Идентификатор клиента": [1.0, None, "abc"]})
    out = format_text_id_columns(frame, config)
    assert out.loc[0, "Идентификатор клиента"] == "1"
    assert pd.isna(out.loc[1, "Идентификатор клиента"]) or out.loc[
        1, "Идентификатор клиента"
    ] is None
    assert out.loc[2, "Идентификатор клиента"] == "abc"
